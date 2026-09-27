"""Печать (фаза 6b): реквизиты с наследованием и правами, формы XLSX/DOCX/HTML(PDF),
шаблоны — загрузка, песочница, откат, аудит."""

import io

import pytest
from conftest import COURSE, FACULTY, SCHEDULE_ID, UNIT, ClientFactory, FakeUpstreams
from docx import Document as Docx
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from dutyflow_common.audit import AuditLog

REQ = {
    "approver_position": "Начальник факультета",
    "approver_rank": "полковник",
    "approver_name": "И. И. Иванов",
    "compiler_name": "П. П. Петров",
}


def weasyprint_available() -> bool:
    try:
        import weasyprint  # noqa: F401
    except OSError:
        return False
    return True


async def test_requisites_inherit_and_rights(client_for: ClientFactory) -> None:
    fac_admin = client_for("unit_admin", unit_id=FACULTY)
    r = await fac_admin.put(f"/document-settings/{FACULTY}", json=REQ)
    assert r.status_code == 200, r.text
    assert r.json()["own"]["approver_name"] == "И. И. Иванов"

    # У курса своих нет — действуют реквизиты факультета
    course = (await fac_admin.get(f"/document-settings/{COURSE}")).json()
    assert course["own"] is None
    assert course["inherited_from"] == "Факультет 1"
    assert course["effective"]["approver_position"] == "Начальник факультета"
    assert course["can_edit"] is True

    # Оператор и наблюдатель видят, но не меняют; вышестоящее — не своё
    operator = client_for("operator", unit_id=COURSE)
    assert (await operator.get(f"/document-settings/{COURSE}")).json()["can_edit"] is False
    assert (await operator.put(f"/document-settings/{COURSE}", json=REQ)).status_code == 403
    assert (await fac_admin.put(f"/document-settings/{UNIT}", json=REQ)).status_code == 403

    # Свои реквизиты курса, затем удаление — снова наследуются
    own = await fac_admin.put(
        f"/document-settings/{COURSE}", json={"approver_name": "С. С. Сидоров"}
    )
    assert own.json()["effective"]["approver_name"] == "С. С. Сидоров"
    stale = await fac_admin.put(
        f"/document-settings/{COURSE}", json={"approver_name": "X", "version": 99}
    )
    assert stale.status_code == 409
    cleared = (await fac_admin.delete(f"/document-settings/{COURSE}")).json()
    assert cleared["inherited_from"] == "Факультет 1"


async def test_schedule_forms(client_for: ClientFactory, upstreams: FakeUpstreams) -> None:
    admin = client_for("unit_admin", unit_id=FACULTY)
    await admin.put(f"/document-settings/{FACULTY}", json=REQ)

    r = await admin.get(f"/print/schedules/{SCHEDULE_ID}", params={"format": "xlsx"})
    assert r.status_code == 200, r.text
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws is not None
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "УТВЕРЖДАЮ" in text
    assert "полковник ____________ И. И. Иванов" in text
    assert "рядовой Алексеев И. П." in text
    assert "Группа 111" in text  # ячейка передана дочернему
    assert "не назначено: 1" in text

    html = (
        await admin.get("/print/html/schedule_month", params={"schedule_id": str(SCHEDULE_ID)})
    ).text
    assert "рядовой Алексеев И. П." in html
    assert "ПРОЕКТ" not in html
    upstreams.status = "draft"
    html = (
        await admin.get("/print/html/schedule_month", params={"schedule_id": str(SCHEDULE_ID)})
    ).text
    assert "ПРОЕКТ" in html
    assert "УТВЕРЖДАЮ" not in html  # open-questions №58


async def test_daily_forms(client_for: ClientFactory) -> None:
    admin = client_for("unit_admin", unit_id=FACULTY)
    await admin.put(f"/document-settings/{FACULTY}", json=REQ)
    params = {"unit_id": str(COURSE), "date": "2026-11-05"}

    r = await admin.get("/print/daily", params={**params, "format": "docx"})
    assert r.status_code == 200, r.text
    doc = Docx(io.BytesIO(r.content))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Алексеев Иван Петрович" in cells
    assert "не назначен" in cells
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "5 ноября 2026 г." in body
    # Время заступления — местное (UTC+3)
    assert any("с 18:00 05.11 до 18:00 06.11" in c for c in cells)

    order = await admin.get(
        "/print/daily", params={**params, "form": "daily_order", "format": "docx"}
    )
    text = "\n".join(p.text for p in Docx(io.BytesIO(order.content)).paragraphs)
    assert "ПРИКАЗЫВАЮ" in text
    assert "рядовой Алексеев Иван Петрович (Группа 111)" in text

    html = (await admin.get("/print/html/daily_order", params=params)).text
    assert "ПРИКАЗЫВАЮ" in html
    assert (await admin.get("/print/daily", params={**params, "format": "xlsx"})).status_code == 422


async def test_load_report(client_for: ClientFactory, upstreams: FakeUpstreams) -> None:
    admin = client_for("unit_admin", unit_id=FACULTY)
    await admin.put(f"/document-settings/{FACULTY}", json={**REQ, "compiler_name": "П. П. Петров"})
    params = {"unit_id": str(COURSE), "date_from": "2026-11-01", "date_to": "2026-11-30"}
    r = await admin.get("/print/load-report", params={**params, "format": "xlsx"})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    people = wb["Люди"]
    rows = [[c.value for c in row] for row in people.iter_rows(min_row=4)]
    assert rows[0][:5] == [1, "Алексеев И. П.", 3, 4, 4.0]
    summary = "\n".join(str(c.value) for row in wb["Сводка"].iter_rows() for c in row if c.value)
    assert "Нагрузка на человека" in summary
    assert "0,96" in summary  # индекс Джайна
    # Запросы к analytics — от имени оператора, с теми же параметрами
    assert any(c.startswith("analytics /metrics/people") for c in upstreams.calls)

    html = (await admin.get("/print/html/load_report", params=params)).text
    assert "01.11.2026 — 30.11.2026" in html
    assert "Составил" in html
    assert "П. П. Петров" in html


async def test_templates_sandbox_and_audit(
    client_for: ClientFactory, app: object, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    admin = client_for("unit_admin", unit_id=FACULTY)
    root = client_for("superadmin", username="root")
    listed = (await admin.get("/templates")).json()
    assert [t["form"] for t in listed] == [
        "schedule_month",
        "daily_roster",
        "daily_order",
        "load_report",
    ]
    assert not any(t["custom"] for t in listed)

    body = "<h1>Свой шаблон {{ unit.name }}</h1>{% for d in duties %}{{ d.name }}{% endfor %}"
    assert (await admin.put("/templates/daily_roster", json={"body": body})).status_code == 403

    broken = await root.put("/templates/daily_roster", json={"body": "{% for x in %}"})
    assert broken.status_code == 422
    assert "строка 1" in broken.json()["message"]
    escape = await root.put(
        "/templates/daily_roster", json={"body": "{{ ''.__class__.__mro__[1].__subclasses__() }}"}
    )
    assert escape.status_code == 422
    assert "недоступным" in escape.json()["message"]

    saved = await root.put("/templates/daily_roster", json={"body": body, "comment": "Проба"})
    assert saved.status_code == 200, saved.text
    assert saved.json()["custom"] is True
    html = (
        await admin.get(
            "/print/html/daily_roster", params={"unit_id": str(COURSE), "date": "2026-11-05"}
        )
    ).text
    assert html == "<h1>Свой шаблон 1 курс</h1>Наряд по курсу"

    reset = (await root.post("/templates/daily_roster/reset")).json()
    assert reset["custom"] is False
    assert reset["body"] == reset["builtin"]
    async with sessionmaker() as s:
        actions = sorted(await s.scalars(select(AuditLog.action)))
    assert actions == ["print_template.reset", "print_template.upload"]


@pytest.mark.skipif(not weasyprint_available(), reason="нет системных библиотек WeasyPrint")
async def test_pdf(client_for: ClientFactory) -> None:
    admin = client_for("unit_admin", unit_id=FACULTY)
    r = await admin.get(f"/print/schedules/{SCHEDULE_ID}")
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")
    r = await admin.get("/print/daily", params={"unit_id": str(COURSE), "date": "2026-11-05"})
    assert r.content.startswith(b"%PDF")
    root = client_for("superadmin", username="root")
    preview = await root.post("/templates/schedule_month/preview", json={})
    assert preview.content.startswith(b"%PDF")
