"""Импорт в documents: шаблон, разбор xlsx/csv, предпросмотр, отчёт, применение, права."""

import datetime as dt
import io

from conftest import ClientFactory, FakePersonnel
from openpyxl import Workbook, load_workbook

from documents.sheets import DATA_SHEET, LISTS_SHEET


def xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = DATA_SHEET
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_template_has_lists_and_passes_token(
    client_for: ClientFactory, personnel: FakePersonnel
) -> None:
    c = client_for("unit_admin")
    r = await c.get("/imports/templates/people")
    assert r.status_code == 200
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb[DATA_SHEET]
    assert [c.value for c in ws[1]] == [
        "Личный номер",
        "Фамилия *",
        "Имя *",
        "Подразделение *",
        "Дата рождения",
    ]
    formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any(LISTS_SHEET in (f or "") for f in formulas)
    assert wb[LISTS_SHEET].sheet_state == "hidden"
    # Вызов personnel — токеном оператора
    method, path, _, auth = personnel.calls[0]
    assert (method, path) == ("GET", "/imports/people/template")
    assert auth == c.headers["Authorization"]


async def test_upload_preview_report_and_apply(
    client_for: ClientFactory, personnel: FakePersonnel
) -> None:
    c = client_for("unit_admin")
    content = xlsx(
        [
            ["Фамилия*", " имя ", "Подразделение", "Дата рождения", "Лишний"],
            ["Иванов", "Иван", "Факультет", dt.datetime(2005, 3, 1), "x"],
            [None, None, None, None, None],
            [None, "Пётр", "Факультет / 1 курс", "01.02.2004", None],
        ]
    )
    r = await c.post(
        "/imports/people", files={"file": ("люди.xlsx", content, "application/octet-stream")}
    )
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["summary"] == {"create": 1, "update": 0, "unchanged": 0, "error": 1}
    assert job["total"] == 2
    assert job["notes"] == ["Столбец «Лишний» не из шаблона — пропущен"]
    sent = personnel.calls[-1][2]
    assert sent is not None
    assert sent["rows"][0] == {
        "row": 2,
        "values": {
            "last_name": "Иванов",
            "first_name": "Иван",
            "unit": "Факультет",
            "attr:birth": "2005-03-01",
        },
    }
    assert sent["rows"][1]["row"] == 4  # пустая строка пропущена, номер — как в файле

    errors = (await c.get(f"/imports/{job['id']}/rows", params={"action": "error"})).json()
    assert errors["total"] == 1
    assert errors["items"][0]["errors"][0]["column"] == "last_name"

    report = await c.get(f"/imports/{job['id']}/report")
    ws = load_workbook(io.BytesIO(report.content))[DATA_SHEET]
    assert ws.cell(row=1, column=6).value == "Результат"
    assert ws.cell(row=3, column=6).value == "Ошибка"
    assert "Фамилия обязательна" in ws.cell(row=3, column=7).value
    assert ws.cell(row=3, column=2).fill.fgColor.rgb.endswith("FFC7CE")

    # Чужая задача не видна
    other = client_for("unit_admin", username="other")
    assert (await other.get(f"/imports/{job['id']}")).status_code == 404

    # «Всё или ничего»: с ошибкой без skip_invalid — 422 от personnel, задача остаётся
    r = await c.post(f"/imports/{job['id']}/apply", json={})
    assert r.status_code == 422
    r = await c.post(f"/imports/{job['id']}/apply", json={"skip_invalid": True})
    assert r.status_code == 200
    assert r.json()["status"] == "applied"
    assert r.json()["applied_rows"] == 1
    assert (await c.post(f"/imports/{job['id']}/apply", json={})).status_code == 422
    listed = (await c.get("/imports")).json()
    assert [j["status"] for j in listed] == ["applied"]


async def test_stale_apply_rechecks(client_for: ClientFactory, personnel: FakePersonnel) -> None:
    c = client_for("unit_admin")
    csv_text = "Фамилия;Имя;Подразделение\nСидоров;Олег;Факультет\n"
    r = await c.post(
        "/imports/people",
        files={"file": ("люди.csv", csv_text.encode("cp1251"), "text/csv")},
    )
    assert r.status_code == 201, r.text
    job = r.json()
    personnel.stale = True
    r = await c.post(f"/imports/{job['id']}/apply", json={})
    assert r.status_code == 409
    assert r.json()["code"] == "import_stale"
    # Перепроверка прошла автоматически: последним был dry-run
    assert personnel.calls[-1][2]["dry_run"] is True  # type: ignore[index]
    personnel.stale = False
    assert (await c.post(f"/imports/{job['id']}/apply", json={})).status_code == 200


async def test_bad_files(client_for: ClientFactory) -> None:
    c = client_for("unit_admin")

    async def upload(name: str, content: bytes) -> str:
        r = await c.post("/imports/people", files={"file": (name, content, "text/plain")})
        assert r.status_code == 422, r.text
        message: str = r.json()["message"]
        return message

    assert "xlsx и .csv" in await upload("a.txt", b"x")
    assert "прочитать" in await upload("a.xlsx", b"not a zip")
    assert "пустой" in await upload("a.csv", b"")
    assert "обязательных столбцов" in await upload("a.csv", "Фамилия,Имя\nА,Б\n".encode())
    assert "нет строк" in await upload("a.csv", "Фамилия,Имя,Подразделение\n,,\n".encode())
    assert "ни одного столбца" in await upload("a.csv", b"x,y\n1,2\n")
