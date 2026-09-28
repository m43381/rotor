"""Сводный журнал аудита (фаза 7b, open-questions №64): события, права, фильтры, выгрузка,
заполнение из сервисов."""

import datetime as dt
import io
import uuid
from typing import Any

from conftest import ClientFactory, Org
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from analytics.journal import handle_audit_event, rebuild_journal, uuid7_time
from analytics.models import AuditEntry
from dutyflow_common.events import Event
from dutyflow_common.ids import uuid7


def entry(
    *,
    unit: uuid.UUID | None,
    action: str = "person.update",
    actor: str = "Иванов Пётр",
    service: str = "personnel",
    when: dt.datetime | None = None,
    **extra: Any,
) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "audit_id": str(uuid7()),
        "service": service,
        "actor_id": str(uuid.uuid4()),
        "actor_name": actor,
        "actor_unit_id": None,
        "action": action,
        "entity_type": action.split(".")[0],
        "entity_id": str(uuid.uuid4()),
        "scope_unit_id": str(unit) if unit else None,
        "before": {"rank_id": None},
        "after": {"rank_id": "x"},
        "comment": None,
        "request_id": "r1",
        **extra,
    }
    if when is not None:
        payload["occurred_at"] = when.isoformat()
    return payload


async def emit(sessionmaker: async_sessionmaker[AsyncSession], payload: dict[str, Any]) -> None:
    async with sessionmaker() as session, session.begin():
        await handle_audit_event(
            session,
            Event(uuid7(), "audit.recorded", uuid.UUID(payload["entity_id"]), payload),
        )


async def count(sessionmaker: async_sessionmaker[AsyncSession]) -> int:
    async with sessionmaker() as s:
        return int(await s.scalar(select(func.count()).select_from(AuditEntry)) or 0)


async def test_events_scope_and_filters(
    client_for: ClientFactory, org: Org, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    now = dt.datetime.now(dt.UTC)
    items = [
        entry(unit=org.course_a1, when=now),
        entry(
            unit=org.fac_a,
            action="assignment.create",
            actor="Петров Сергей",
            service="scheduling",
            when=now,
        ),
        entry(unit=org.fac_b, when=now),
        entry(unit=None, action="rank.update", service="org", when=now),  # справочник
        entry(unit=org.course_a1, when=now - dt.timedelta(days=40)),
    ]
    for p in items:
        await emit(sessionmaker, p)
    await emit(sessionmaker, items[0])  # повтор события
    assert await count(sessionmaker) == 5

    op = client_for(org.fac_a, "operator")
    page = (await op.get("/audit")).json()
    assert page["total"] == 3  # своё поддерево, без справочников и чужого факультета
    assert {i["unit_name"] for i in page["items"]} == {"Курс A1", "Факультет A"}
    root = client_for(org.root, "superadmin")
    assert (await root.get("/audit")).json()["total"] == 5
    viewer = client_for(org.fac_a, "viewer")
    assert (await viewer.get("/audit")).status_code == 403

    today = dt.date.today()
    recent = await op.get(
        "/audit", params={"date_from": str(today - dt.timedelta(days=7)), "date_to": str(today)}
    )
    assert recent.json()["total"] == 2
    by_actor = (await op.get("/audit", params={"actor": "петров"})).json()
    assert [i["action"] for i in by_actor["items"]] == ["assignment.create"]
    by_prefix = (await op.get("/audit", params={"action": "person."})).json()
    assert by_prefix["total"] == 2
    by_unit = (await op.get("/audit", params={"unit_id": str(org.course_a1)})).json()
    assert by_unit["total"] == 2
    assert (await op.get("/audit", params={"unit_id": str(org.fac_b)})).status_code == 404

    facets = (await root.get("/audit/facets")).json()
    assert facets["service"] == ["org", "personnel", "scheduling"]
    assert "rank" in facets["entity_type"]

    export = await op.get("/audit/export", params={"action": "person."})
    assert export.status_code == 200
    ws = load_workbook(io.BytesIO(export.content)).active
    assert ws is not None
    assert ws.max_row == 3  # заголовок + 2 записи
    assert ws.cell(row=2, column=7).value == "rank_id: None"


async def test_old_events_and_rebuild(
    org: Org, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    # Событие без времени (до фазы 7b) — время берётся из uuid7
    old = entry(unit=org.fac_a)
    await emit(sessionmaker, old)
    async with sessionmaker() as s:
        stored = await s.get(AuditEntry, uuid.UUID(old["audit_id"]))
        assert stored is not None
        assert stored.occurred_at == uuid7_time(uuid.UUID(old["audit_id"]))
        assert abs((stored.occurred_at - dt.datetime.now(dt.UTC)).total_seconds()) < 60

    class Source:
        def __init__(self, rows: list[dict[str, Any]]) -> None:
            self.rows = rows

        async def get(self, path: str) -> list[dict[str, Any]]:
            after = path.split("after=")[1] if "after=" in path else None
            return [r for r in self.rows if after is None or r["audit_id"] > after]

    class Broken:
        async def get(self, path: str) -> list[dict[str, Any]]:
            raise RuntimeError("недоступен")

    rows = [entry(unit=org.fac_a, service="org") for _ in range(3)] + [old]
    await rebuild_journal(
        sessionmaker,
        {"org": Source(rows), "documents": Broken()},  # type: ignore[dict-item]
    )
    assert await count(sessionmaker) == 4  # уже известная запись не дублируется
