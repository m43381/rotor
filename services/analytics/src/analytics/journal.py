"""Сводный журнал аудита (ADR-0010, фаза 7b, open-questions №64).

Записи всех сервисов приходят событием `audit.recorded` (поток `events:audit`); журнал
только дополняется — повторная доставка ничего не меняет. До появления журнала (или после
восстановления БД) он заполняется выгрузкой `GET /internal/audit` каждого сервиса.

Кто видит: суперадминистратор — всё, включая записи без подразделения (справочники,
шаблоны документов); администратор подразделения и оператор — записи своего поддерева;
наблюдатель — нет. Хранится бессрочно, выгружается в xlsx.
"""

import datetime as dt
import io
import logging
import uuid
from collections.abc import Iterable, Mapping
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import Select, func, or_, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from analytics.models import AuditEntry
from dutyflow_common.context import Operator
from dutyflow_common.errors import ForbiddenError, NotFoundError, ValidationFailedError
from dutyflow_common.events import Event
from dutyflow_common.internal import InternalClient
from dutyflow_common.ltree import is_descendant_or_self
from dutyflow_common.policy import Policy, default_policy
from dutyflow_common.projections import UnitProjection, unit_path
from dutyflow_common.scope import Scope, in_scope, scope_clause

log = logging.getLogger(__name__)
PAGE = 2_000
EXPORT_LIMIT = 50_000


def _uuid(value: Any) -> uuid.UUID | None:
    return uuid.UUID(str(value)) if value else None


def uuid7_time(value: uuid.UUID) -> dt.datetime:
    """Время создания из uuid7: первые 48 бит — миллисекунды Unix."""
    return dt.datetime.fromtimestamp((value.int >> 80) / 1000, dt.UTC)


def row(p: Mapping[str, Any]) -> dict[str, Any]:
    audit_id = uuid.UUID(str(p["audit_id"]))
    occurred = p.get("occurred_at")
    return {
        "id": audit_id,
        "service": str(p.get("service") or "unknown"),
        # События до фазы 7b не несут времени — берём его из uuid7 записи
        "occurred_at": dt.datetime.fromisoformat(str(occurred))
        if occurred
        else uuid7_time(audit_id),
        "actor_id": str(p["actor_id"]),
        "actor_name": str(p["actor_name"]),
        "actor_unit_id": _uuid(p.get("actor_unit_id")),
        "action": str(p["action"]),
        "entity_type": str(p["entity_type"]),
        "entity_id": uuid.UUID(str(p["entity_id"])),
        "scope_unit_id": _uuid(p.get("scope_unit_id")),
        "before": p.get("before"),
        "after": p.get("after"),
        "comment": p.get("comment"),
        "request_id": str(p.get("request_id") or ""),
    }


async def add_entries(session: AsyncSession, payloads: Iterable[Mapping[str, Any]]) -> int:
    rows = [row(p) for p in payloads]
    for i in range(0, len(rows), 1_000):
        await session.execute(
            insert(AuditEntry).values(rows[i : i + 1_000]).on_conflict_do_nothing()
        )
    return len(rows)


async def handle_audit_event(session: AsyncSession, event: Event) -> None:
    if event.type == "audit.recorded":
        await add_entries(session, [event.payload])


async def rebuild_journal(
    sessionmaker: async_sessionmaker[AsyncSession], services: Mapping[str, InternalClient]
) -> int:
    """Дополняет журнал выгрузкой из сервисов (записи, которых нет, — добавляются)."""
    total = 0
    for name, client in services.items():
        after: str | None = None
        while True:
            query = f"?limit={PAGE}" + (f"&after={after}" if after else "")
            try:
                page = await client.get(f"/internal/audit{query}")
            except Exception as exc:  # сервис недоступен — остальные всё равно заполняются
                log.warning("Журнал %s не выгружен: %s", name, exc)
                break
            if not page:
                break
            async with sessionmaker() as session, session.begin():
                total += await add_entries(session, page)
            after = page[-1]["audit_id"]
            if len(page) < PAGE:
                break
    log.info("Сводный журнал аудита дополнен из сервисов: %d записей просмотрено", total)
    return total


class JournalService:
    def __init__(
        self,
        session: AsyncSession,
        operator: Operator,
        tz: ZoneInfo,
        policy: Policy = default_policy,
    ) -> None:
        self.session = session
        self.operator = operator
        self.tz = tz
        self.policy = policy

    async def _query(
        self,
        *,
        date_from: dt.date | None,
        date_to: dt.date | None,
        unit_id: uuid.UUID | None,
        actor: str | None,
        entity_type: str | None,
        entity_id: uuid.UUID | None,
        action: str | None,
        service: str | None,
    ) -> Select[AuditEntry, str | None]:
        scope = self.policy.scope_for(self.operator.roles, "audit_journal", "read")
        if scope is Scope.NONE:
            raise ForbiddenError("Журнал аудита недоступен для этой роли")
        stmt = select(AuditEntry, UnitProjection.name).outerjoin(
            UnitProjection, UnitProjection.unit_id == AuditEntry.scope_unit_id
        )
        if scope is not Scope.ALL:
            own = await unit_path(self.session, self.operator.unit_id)
            stmt = stmt.where(scope_clause(UnitProjection.path, own, scope))
        if unit_id is not None:
            unit = await self.session.get(UnitProjection, unit_id)
            own = await unit_path(self.session, self.operator.unit_id)
            if unit is None or not in_scope(unit.path, own, scope):
                raise NotFoundError("Подразделение не найдено")
            stmt = stmt.where(is_descendant_or_self(UnitProjection.path, unit.path))
        if date_from and date_to and date_to < date_from:
            raise ValidationFailedError("Конец периода раньше начала")
        if date_from:
            stmt = stmt.where(
                AuditEntry.occurred_at >= dt.datetime.combine(date_from, dt.time(), self.tz)
            )
        if date_to:
            stmt = stmt.where(
                AuditEntry.occurred_at
                < dt.datetime.combine(date_to + dt.timedelta(days=1), dt.time(), self.tz)
            )
        if actor and actor.strip():
            needle = f"%{actor.strip().lower()}%"
            stmt = stmt.where(
                or_(func.lower(AuditEntry.actor_name).like(needle), AuditEntry.actor_id == actor)
            )
        if entity_type:
            stmt = stmt.where(AuditEntry.entity_type == entity_type)
        if entity_id:
            stmt = stmt.where(AuditEntry.entity_id == entity_id)
        if action:
            stmt = stmt.where(AuditEntry.action.startswith(action))
        if service:
            stmt = stmt.where(AuditEntry.service == service)
        return stmt

    @staticmethod
    def _out(e: AuditEntry, unit_name: str | None) -> dict[str, Any]:
        return {
            "id": e.id,
            "service": e.service,
            "occurred_at": e.occurred_at,
            "actor_id": e.actor_id,
            "actor_name": e.actor_name,
            "action": e.action,
            "entity_type": e.entity_type,
            "entity_id": e.entity_id,
            "unit_id": e.scope_unit_id,
            "unit_name": unit_name,
            "before": e.before,
            "after": e.after,
            "comment": e.comment,
        }

    async def list_entries(self, *, limit: int, offset: int, **filters: Any) -> dict[str, Any]:
        stmt = await self._query(**filters)
        total = await self.session.scalar(select(func.count()).select_from(stmt.subquery()))
        rows = await self.session.execute(
            stmt.order_by(AuditEntry.occurred_at.desc(), AuditEntry.id.desc())
            .limit(limit)
            .offset(offset)
        )
        return {
            "items": [self._out(e, name) for e, name in rows],
            "total": total or 0,
            "limit": limit,
            "offset": offset,
        }

    async def facets(self) -> dict[str, list[str]]:
        """Значения фильтров: виды объектов, действия, сервисы (в пределах scope)."""
        stmt = await self._query(
            date_from=None, date_to=None, unit_id=None, actor=None,
            entity_type=None, entity_id=None, action=None, service=None,
        )  # fmt: skip
        sub = stmt.subquery()
        result: dict[str, list[str]] = {}
        for key in ("entity_type", "action", "service"):
            values: list[str] = list(
                await self.session.scalars(select(sub.c[key]).distinct().order_by(sub.c[key]))
            )
            result[key] = values
        return result

    async def export(self, **filters: Any) -> bytes:
        stmt = await self._query(**filters)
        rows = (
            await self.session.execute(
                stmt.order_by(AuditEntry.occurred_at.desc()).limit(EXPORT_LIMIT)
            )
        ).all()
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "Журнал"
        header = ["Когда", "Кто", "Действие", "Объект", "Id объекта", "Подразделение",
                  "Было", "Стало", "Комментарий", "Сервис"]  # fmt: skip
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for e, unit_name in rows:
            ws.append(
                [
                    e.occurred_at.astimezone(self.tz).strftime("%d.%m.%Y %H:%M:%S"),
                    e.actor_name,
                    e.action,
                    e.entity_type,
                    str(e.entity_id),
                    unit_name or "",
                    _flat(e.before),
                    _flat(e.after),
                    e.comment or "",
                    e.service,
                ]
            )
        for letter, width in zip(
            "ABCDEFGHIJ", (19, 24, 26, 16, 38, 24, 50, 50, 30, 14), strict=True
        ):
            ws.column_dimensions[letter].width = width
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:J{ws.max_row}"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def _flat(values: dict[str, Any] | None) -> str:
    if not values:
        return ""
    return "\n".join(f"{k}: {v}" for k, v in values.items())
